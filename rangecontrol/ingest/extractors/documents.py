"""PDF, Word, PowerPoint, and spreadsheet extraction.

Every extractor raises ValueError when it produces no usable text. Callers
record that as an unreadable document rather than aborting startup.

Pictures embedded in PDF, DOCX, and PPTX files are described through the
vision path (see ``embedded``); a diagram inside a brief is range material.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from pathlib import Path

from rangecontrol.ingest.extractors.embedded import (
    EMBEDDED_IMAGES_HEADER,
    cap_note,
    describe_embedded,
)
from rangecontrol.llm.base import Provider

logger = logging.getLogger(__name__)

# Layout-mode PDF text keeps table columns apart with runs of spaces, which
# is what makes an asset inventory readable as rows. Two spaces is enough to
# mark a column boundary; longer runs only cost tokens.
_SPACE_RUNS = re.compile(r" {2,}")

# An encrypted .docx/.pptx/.xlsx is not a zip at all but an OLE compound
# file wrapping an EncryptedPackage stream. The libraries report it as "not
# a package", which tells an operator nothing about the fix.
_OLE_MAGIC = bytes.fromhex("d0cf11e0a1b11ae1")
ENCRYPTED_OFFICE_TEXT = (
    "this file is password-protected (encrypted). Remove the password in "
    "Office (File > Info > Protect Document) and save it again."
)

# A PDF protected with Microsoft Information Rights Management keeps the
# real document in an encrypted stream and shows one placeholder page
# saying so. That page used to be extracted as if it were the content.
_IRM_MARKERS = ("this pdf document has been protected", "protected by microsoft office")
IRM_PDF_TEXT = (
    "this PDF is protected by Microsoft Information Rights Management; only "
    "a placeholder page is readable. Export an unprotected copy from Office."
)
ENCRYPTED_PDF_TEXT = (
    "this PDF is password-protected. Remove the password and save it again."
)


def _reject_encrypted_office(path: Path) -> None:
    with path.open("rb") as handle:
        head = handle.read(len(_OLE_MAGIC))
    if head == _OLE_MAGIC:
        raise ValueError(ENCRYPTED_OFFICE_TEXT)


# --- PDF -------------------------------------------------------------------


def _page_text(page) -> str:
    """Layout-preserving text with a plain fallback.

    pypdf's layout mode reconstructs columns, which is what keeps a table of
    hosts and addresses aligned instead of interleaved. It can fail on
    unusual fonts or return nothing where the plain mode still succeeds, so
    the plain extraction is the safety net.
    """
    try:
        text = page.extract_text(extraction_mode="layout") or ""
    except Exception:  # noqa: BLE001 - layout mode is best-effort
        text = ""
    if not text.strip():
        text = page.extract_text() or ""
    lines = [_SPACE_RUNS.sub("  ", line.rstrip()) for line in text.splitlines()]
    return "\n".join(lines).strip()


def _page_images(page, number: int) -> list[tuple[str, bytes]]:
    found: list[tuple[str, bytes]] = []
    try:
        for index, item in enumerate(page.images, start=1):
            found.append((f"page {number} image {index}", item.data))
    except Exception as exc:  # noqa: BLE001 - an odd image stream must not lose the page
        logger.warning("page %s: could not read embedded images: %s", number, exc)
    return found


def extract_pdf(path: Path, provider: Provider) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            # An owner-password-only PDF opens with the empty user password.
            try:
                opened = reader.decrypt("")
            except Exception:  # noqa: BLE001 - treat as locked
                opened = 0
            if not opened:
                raise ValueError(ENCRYPTED_PDF_TEXT)
        pages = list(reader.pages)
        texts = [_page_text(page) for page in pages]
    except ValueError:
        raise
    except Exception as exc:  # noqa: BLE001 - any parse failure is "unreadable"
        raise ValueError(f"could not parse PDF: {type(exc).__name__}") from exc

    if len(pages) == 1 and any(marker in texts[0].lower() for marker in _IRM_MARKERS):
        raise ValueError(IRM_PDF_TEXT)

    images: list[tuple[str, bytes]] = []
    for number, page in enumerate(pages, start=1):
        images.extend(_page_images(page, number))
    descriptions = describe_embedded(images, provider)

    sections = [
        f"--- page {number} ---\n{body}"
        for number, body in enumerate(texts, start=1)
        if body
    ]
    described = [entry for entry in descriptions if entry]
    if described:
        sections.append(f"{EMBEDDED_IMAGES_HEADER}\n" + "\n\n".join(described))
    note = cap_note(images)
    if note:
        sections.append(note)

    if not sections:
        raise ValueError(
            "PDF contained no extractable text and no describable images"
        )
    return "\n\n".join(sections)


# --- Word --------------------------------------------------------------------


def _image_parts(part) -> list:
    return [
        related
        for related in part.related_parts.values()
        if str(getattr(related, "content_type", "") or "").startswith("image/")
    ]


def _docx_images(document) -> list[tuple[str, bytes]]:
    """Pictures in the body, then in headers and footers.

    Headers and footers are separate parts with their own relationships; a
    diagram used as a header graphic would otherwise never be reached.
    """
    found: list[tuple[str, bytes]] = []
    seen: set[str] = set()
    try:
        parts = [document.part]
        for section in document.sections:
            parts.extend(block.part for block in (section.header, section.footer))
        for part in parts:
            for related in _image_parts(part):
                name = str(getattr(related, "partname", "") or "")
                if name in seen:
                    continue
                seen.add(name)
                found.append((f"image {len(found) + 1}: {name.rsplit('/', 1)[-1]}", related.blob))
    except Exception as exc:  # noqa: BLE001 - pictures are a bonus, text is the document
        logger.warning("could not read embedded images: %s", exc)
    return found


def _docx_header_footer_lines(document) -> list[str]:
    lines: list[str] = []
    seen: set[str] = set()
    try:
        for section in document.sections:
            for block in (section.header, section.footer):
                for paragraph in block.paragraphs:
                    text = paragraph.text.strip()
                    if text and text not in seen:
                        seen.add(text)
                        lines.append(text)
    except Exception:  # noqa: BLE001 - headers are optional context
        return lines
    return lines


def extract_docx(path: Path, provider: Provider) -> str:
    import docx

    _reject_encrypted_office(path)
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"could not parse DOCX: {type(exc).__name__}") from exc

    lines = _docx_header_footer_lines(document)
    lines.extend(p.text.strip() for p in document.paragraphs if p.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    images = _docx_images(document)
    described = [entry for entry in describe_embedded(images, provider) if entry]
    if described:
        lines.append(EMBEDDED_IMAGES_HEADER)
        lines.extend(described)
    note = cap_note(images)
    if note:
        lines.append(note)

    if not lines:
        raise ValueError("DOCX contained no text and no describable images")
    return "\n".join(lines)


# --- PowerPoint ----------------------------------------------------------------


def _iter_shapes(shapes):
    """Walk shapes depth-first, descending into groups."""
    for shape in shapes:
        if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _shape_picture(shape) -> bytes | None:
    """The picture bytes of a shape, or None when it is not a picture.

    Covers both free-floating pictures and picture placeholders that a
    templated layout fills: the latter report a PLACEHOLDER shape type, so
    matching on the type alone misses every diagram dropped into a layout.
    """
    if not hasattr(shape, "image"):
        return None
    try:
        return shape.image.blob
    except Exception:  # noqa: BLE001 - an empty placeholder raises; not a picture
        return None


def _slide_lines(slide, number: int, images: list[tuple[str, bytes]]) -> list[str]:
    lines: list[str] = []
    picture_index = 0
    for shape in _iter_shapes(slide.shapes):
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                lines.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        blob = _shape_picture(shape)
        if blob is not None:
            picture_index += 1
            images.append((f"slide {number} image {picture_index}", blob))
    if getattr(slide, "has_notes_slide", False) and slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            lines.append(f"[speaker notes] {notes}")
    return lines


def extract_pptx(path: Path, provider: Provider) -> str:
    from pptx import Presentation

    _reject_encrypted_office(path)
    try:
        presentation = Presentation(str(path))
        slides = list(presentation.slides)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"could not parse PPTX: {type(exc).__name__}") from exc

    images: list[tuple[str, bytes]] = []
    per_slide: list[tuple[int, list[str], int]] = []
    for number, slide in enumerate(slides, start=1):
        before = len(images)
        lines = _slide_lines(slide, number, images)
        per_slide.append((number, lines, before))

    descriptions = describe_embedded(images, provider)

    sections: list[str] = []
    for index, (number, lines, start) in enumerate(per_slide):
        end = per_slide[index + 1][2] if index + 1 < len(per_slide) else len(images)
        pictures = [entry for entry in descriptions[start:end] if entry]
        body = lines + pictures
        if body:
            sections.append(f"--- slide {number} ---\n" + "\n".join(body))
    note = cap_note(images)
    if note:
        sections.append(note)

    if not sections:
        raise ValueError("PPTX contained no text and no describable images")
    return "\n\n".join(sections)


# --- Spreadsheets ------------------------------------------------------------


def extract_spreadsheet(path: Path, provider: Provider) -> str:
    if path.suffix.lower() in {".csv", ".tsv"}:
        return _extract_delimited(path)
    return _extract_xlsx(path)


def _extract_delimited(path: Path) -> str:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    raw = path.read_text(encoding="utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
    lines = []
    for row in rows:
        cells = [cell.strip() for cell in row]
        if any(cells):
            lines.append(" | ".join(cells))

    if not lines:
        raise ValueError("spreadsheet contained no rows")
    return "\n".join(lines)


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    _reject_encrypted_office(path)
    try:
        book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"could not parse workbook: {type(exc).__name__}") from exc

    sections: list[str] = []
    try:
        for sheet in book.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    lines.append(" | ".join(cells))
            if lines:
                sections.append(f"--- sheet: {sheet.title} ---\n" + "\n".join(lines))
    except Exception as exc:  # noqa: BLE001 - lazy parse failure mid-iteration
        raise ValueError(f"could not read workbook rows: {type(exc).__name__}") from exc
    finally:
        book.close()

    if not sections:
        raise ValueError("workbook contained no rows")
    return "\n\n".join(sections)
