"""Diagrams inside documents are range material too.

A network diagram pasted into a Word brief, a slide deck, or exported PDF is
exactly the material an adjudicator needs, and text extraction alone drops
it silently.
"""

import pytest

from rangecontrol.ingest.extractors.documents import (
    extract_docx,
    extract_pdf,
    extract_pptx,
)
from rangecontrol.ingest.extractors.embedded import MAX_EMBEDDED_IMAGES
from rangecontrol.llm.base import LLMError
from tests.support.stub_provider import StubProvider


def _png_file(tmp_path, name="diagram.png", size=(160, 120)):
    Image = pytest.importorskip("PIL.Image")
    path = tmp_path / name
    Image.new("RGB", size, (200, 30, 30)).save(path)
    return path


def _pptx(tmp_path, *, pictures=1, tiny=False, name="brief.pptx"):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    picture = _png_file(tmp_path, size=(24, 24) if tiny else (160, 120))
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Network Plan"
    for _ in range(pictures):
        slide.shapes.add_picture(str(picture), Inches(1), Inches(2))
    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    box.text_frame.text = "DC-01 lives at 10.77.0.5"
    table = slide.shapes.add_table(2, 2, Inches(5), Inches(2), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "Host"
    table.cell(0, 1).text = "IP"
    table.cell(1, 0).text = "FW-1"
    table.cell(1, 1).text = "10.77.0.1"
    slide.notes_slide.notes_text_frame.text = "Inject 3 fires at 1400"
    path = tmp_path / name
    prs.save(str(path))
    return path


def _pptx_with_placeholder_picture(tmp_path):
    """A templated 'Picture with Caption' slide reports the picture as a
    PLACEHOLDER shape, not a PICTURE; it must still be described."""
    pptx = pytest.importorskip("pptx")

    prs = pptx.Presentation()
    layout = next(
        layout for layout in prs.slide_layouts
        if any(ph.placeholder_format.type is not None
               and "PICTURE" in str(ph.placeholder_format.type)
               for ph in layout.placeholders)
    )
    slide = prs.slides.add_slide(layout)
    holder = next(
        ph for ph in slide.placeholders if "PICTURE" in str(ph.placeholder_format.type)
    )
    holder.insert_picture(str(_png_file(tmp_path)))
    path = tmp_path / "templated.pptx"
    prs.save(str(path))
    return path


def test_pptx_extracts_text_tables_notes_and_every_picture(tmp_path):
    stub = StubProvider(image_text="Two subnets joined by FW-1.")
    text = extract_pptx(_pptx(tmp_path), stub)
    assert "--- slide 1 ---" in text
    assert "Network Plan" in text
    assert "DC-01 lives at 10.77.0.5" in text
    assert "FW-1 | 10.77.0.1" in text
    assert "Inject 3 fires at 1400" in text
    assert "Two subnets joined by FW-1." in text
    assert len(stub.image_calls) == 1

    placeholder = StubProvider(image_text="Diagram in a placeholder.")
    assert "Diagram in a placeholder." in extract_pptx(
        _pptx_with_placeholder_picture(tmp_path), placeholder
    )


def test_docx_body_and_header_pictures_and_image_only_pdfs_are_described(tmp_path):
    docx = pytest.importorskip("docx")
    Image = pytest.importorskip("PIL.Image")

    body = docx.Document()
    body.add_paragraph("Topology below.")
    body.add_picture(str(_png_file(tmp_path)))
    body_path = tmp_path / "brief.docx"
    body.save(str(body_path))
    text = extract_docx(body_path, StubProvider(image_text="A DMZ behind FW-1."))
    assert "Topology below." in text
    assert "A DMZ behind FW-1." in text

    headed = docx.Document()
    headed.add_paragraph("Body text.")
    headed.sections[0].header.paragraphs[0].add_run().add_picture(str(_png_file(tmp_path)))
    headed_path = tmp_path / "header.docx"
    headed.save(str(headed_path))
    stub = StubProvider(image_text="Header diagram.")
    assert "Header diagram." in extract_docx(headed_path, stub)
    assert len(stub.image_calls) == 1

    scan = tmp_path / "scan.pdf"
    Image.new("RGB", (300, 200), (10, 10, 200)).save(scan)  # an image-only PDF
    pdf_text = extract_pdf(scan, StubProvider(image_text="Diagram: HQ-LAN to DMZ via FW-1."))
    assert "HQ-LAN to DMZ" in pdf_text


def test_icons_overflow_and_failed_descriptions_never_lose_the_document(tmp_path):
    icons = StubProvider(image_text="should not be asked")
    text = extract_pptx(_pptx(tmp_path, tiny=True, name="icons.pptx"), icons)
    assert icons.image_calls == []
    assert "should not be asked" not in text

    many = StubProvider(image_text="desc")
    text = extract_pptx(_pptx(tmp_path, pictures=MAX_EMBEDDED_IMAGES + 3, name="many.pptx"), many)
    assert len(many.image_calls) == MAX_EMBEDDED_IMAGES
    assert "3 more image(s)" in text

    down = StubProvider(error=LLMError("vision down"))
    text = extract_pptx(_pptx(tmp_path, name="down.pptx"), down)
    assert "DC-01 lives at 10.77.0.5" in text
    assert "could not be described" in text
    assert "vision down" in text
